from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from config import get_workspace_settings


REPORT_FIELDS = [
    ("name", "Name"),
    ("school", "School"),
    ("college", "College"),
    ("title", "Title"),
    ("direction", "Direction"),
    ("email", "Email"),
    ("admission_type", "Admission Type"),
    ("match_score", "Match Score"),
    ("match_level", "Match Level"),
    ("match_reason", "Match Reason"),
    ("risk_note", "Risk Note"),
    ("source", "Source"),
]

DDL_FIELDS = [
    ("name", "Title"),
    ("institute", "Institute"),
    ("province", "Province"),
    ("deadline_date", "Deadline"),
    ("days_left", "Days Left"),
    ("status", "Status"),
    ("description", "Description"),
    ("website", "Link"),
]


def _report_dir() -> Path:
    directory = get_workspace_settings().word_reports_dir
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return ", ".join(_cell(item) for item in value)
    return str(value)


def _safe_name(title: str, suffix: str) -> str:
    allowed = []
    for char in title:
        allowed.append(char if char.isalnum() or char in {"-", "_"} else "_")
    return f"{''.join(allowed).strip('_') or 'GradMate_Report'}_{_stamp()}.{suffix}"


def _write_docx(title: str, sections: list[dict[str, Any]], tables: list[dict[str, Any]]) -> Path:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed. Run: pip install python-docx") from exc

    path = _report_dir() / _safe_name(title, "docx")
    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    for section in sections:
        doc.add_heading(section.get("title", "Section"), level=1)
        for line in section.get("lines", []):
            doc.add_paragraph(_cell(line))

    for table_spec in tables:
        rows = table_spec.get("rows", [])
        fields = table_spec.get("fields", [])
        doc.add_heading(table_spec.get("title", "Table"), level=1)
        if not rows:
            doc.add_paragraph("No data.")
            continue
        table = doc.add_table(rows=1, cols=len(fields))
        table.style = "Table Grid"
        for index, (_, label) in enumerate(fields):
            table.rows[0].cells[index].text = label
        for row in rows:
            cells = table.add_row().cells
            for index, (key, _) in enumerate(fields):
                cells[index].text = _cell(row.get(key, ""))

    doc.save(path)
    return path


def _write_xlsx(title: str, sheets: list[dict[str, Any]]) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is not installed. Run: pip install openpyxl") from exc

    path = _report_dir() / _safe_name(title, "xlsx")
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    for sheet_spec in sheets:
        sheet = workbook.create_sheet(title=sheet_spec.get("title", "Sheet")[:31] or "Sheet")
        rows = sheet_spec.get("rows", [])
        fields = sheet_spec.get("fields", [])
        sheet.append([label for _, label in fields])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="263238")
        for row in rows:
            sheet.append([_cell(row.get(key, "")) for key, _ in fields])
        for index, _ in enumerate(fields, start=1):
            column = get_column_letter(index)
            width = max(12, min(42, max(len(_cell(sheet.cell(row=r, column=index).value)) for r in range(1, sheet.max_row + 1)) + 2))
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"

    workbook.save(path)
    return path


def export_supervisors_report(supervisors: list[dict[str, Any]], report_format: str = "docx", title: str = "GradMate Supervisor Report") -> Path:
    rows = list(supervisors)
    if report_format == "xlsx":
        return _write_xlsx(title, [{"title": "Supervisors", "rows": rows, "fields": REPORT_FIELDS}])
    return _write_docx(title, [], [{"title": "Supervisors", "rows": rows, "fields": REPORT_FIELDS}])


def export_ddl_report(items: list[dict[str, Any]], report_format: str = "docx", title: str = "GradMate Admission DDL Report") -> Path:
    rows = list(items)
    if report_format == "xlsx":
        return _write_xlsx(title, [{"title": "DDL", "rows": rows, "fields": DDL_FIELDS}])
    return _write_docx(title, [], [{"title": "DDL", "rows": rows, "fields": DDL_FIELDS}])


def export_knowledge_report(items: list[dict[str, Any]], report_format: str = "docx", title: str = "GradMate Knowledge Base Report") -> Path:
    fields = [("source", "Source"), ("kind", "Kind"), ("text", "Snippet")]
    if report_format == "xlsx":
        return _write_xlsx(title, [{"title": "Knowledge", "rows": items, "fields": fields}])
    return _write_docx(title, [], [{"title": "Knowledge", "rows": items, "fields": fields}])


def export_rag_answer_report(payload: dict[str, Any], report_format: str = "docx") -> Path:
    title = "GradMate RAG Answer Report"
    sections = [
        {"title": "Question", "lines": [payload.get("question", "")]},
        {"title": "Answer", "lines": [payload.get("answer", "")]},
    ]
    sources = payload.get("sources") or []
    fields = [("source", "Source"), ("kind", "Kind"), ("snippet", "Snippet")]
    if report_format == "xlsx":
        rows = [{"field": "Question", "content": payload.get("question", "")}, {"field": "Answer", "content": payload.get("answer", "")}]
        return _write_xlsx(title, [
            {"title": "Answer", "rows": rows, "fields": [("field", "Field"), ("content", "Content")]},
            {"title": "Sources", "rows": sources, "fields": fields},
        ])
    return _write_docx(title, sections, [{"title": "Sources", "rows": sources, "fields": fields}])


def export_questions_report(payload: dict[str, Any], report_format: str = "docx") -> Path:
    title = "GradMate Interview Questions Report"
    groups = payload.get("groups") or []
    rows: list[dict[str, Any]] = []
    for group in groups:
        for item in group.get("items", []):
            rows.append({
                "group": group.get("title", ""),
                "question": item.get("question", ""),
                "intent": item.get("intent", ""),
                "prep": item.get("prep", ""),
            })
    sections = [
        {"title": "Supervisor", "lines": [payload.get("supervisor", "")]},
        {"title": "Summary", "lines": [payload.get("summary", "")]},
    ]
    fields = [("group", "Group"), ("question", "Question"), ("intent", "Intent"), ("prep", "Preparation")]
    if report_format == "xlsx":
        return _write_xlsx(title, [{"title": "Questions", "rows": rows, "fields": fields}])
    return _write_docx(title, sections, [{"title": "Questions", "rows": rows, "fields": fields}])

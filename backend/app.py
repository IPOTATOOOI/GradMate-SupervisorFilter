from __future__ import annotations

import json
import re
import base64
import logging
from datetime import date
import asyncio
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from fastapi import FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from agent_runtime import agent_config, invoke_agent, run_application_assistant
from agent_tools import (
    extract_supervisors_from_text,
    profile_from_dict,
    profile_summary,
    score_supervisors_hybrid,
    supervisor_from_dict,
    supervisor_to_dict,
)
from storage import (
    ensure_data_dir,
    get_knowledge_items,
    get_assistant_history,
    get_profile,
    get_question_history,
    get_rag_history,
    get_supervisors,
    merge_supervisors,
    save_knowledge_items,
    save_profile,
    save_question_history,
    save_rag_history,
    save_supervisors,
)
from config import get_model_settings, get_security_settings, get_server_settings
from ddl_service import get_baoyan_ddl
from email_routes import router as email_router
from practice_routes import router as practice_router
from resume_routes import router as resume_router
from export_tools import (
    export_ddl_report,
    export_knowledge_report,
    export_questions_report,
    export_rag_answer_report,
    export_supervisors_report,
)
from browser_research import browse_school_supervisors
from school_directory import find_school_sites, school_directory_summary


BASE_DIR = Path(__file__).resolve().parent
FRONTEND_DIR = BASE_DIR.parent / "frontend"
FRONTEND_DIST_DIR = FRONTEND_DIR / "dist"
FRONTEND_ASSETS_DIR = FRONTEND_DIST_DIR / "assets"
logger = logging.getLogger(__name__)


class SupervisorUploadRequest(BaseModel):
    text: str = ""
    source: str = "用户上传资料"


class AgentInvokeRequest(BaseModel):
    message: str


class AssistantTaskRequest(BaseModel):
    message: str


class QuestionGenerateRequest(BaseModel):
    supervisor_id: int
    focus: str = ""


class RagAskRequest(BaseModel):
    question: str
    school: str = ""
    college: str = ""
    use_browser: bool = True


class ExportRagAnswerRequest(BaseModel):
    question: str = ""
    answer: str = ""
    sources: list[dict[str, Any]] = []


class ExportQuestionsRequest(BaseModel):
    supervisor: str = ""
    summary: str = ""
    groups: list[dict[str, Any]] = []


class SupervisorBrowseRequest(BaseModel):
    school: str
    major: str = ""
    college: str = ""
    keywords: str = ""
    max_pages: int = 12
    save: bool = True


class SupervisorContactStatusRequest(BaseModel):
    status: str


def _extract_json_object(text: str) -> dict[str, Any]:
    text = (text or "").strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, re.S)
        if not match:
            return {}
        try:
            return json.loads(match.group(0))
        except json.JSONDecodeError:
            return {}


def _is_meaningful_supervisor(item: dict[str, Any]) -> bool:
    name = str(item.get("name") or "").strip()
    if not name or "待确认" in name:
        return False
    if not name or name in {"待确认导师", "寰呯‘璁ゅ甯?"}:
        return False
    return any(str(item.get(field) or "").strip() for field in ["email", "title", "direction", "college"])


def _normalize_llm_upload_records(records: Any, text: str, source: str, start_id: int) -> list[dict[str, Any]]:
    if not isinstance(records, list):
        return []
    items: list[dict[str, Any]] = []
    seen_keys: set[str] = set()
    for record in records:
        if not isinstance(record, dict):
            continue
        email = str(record.get("email") or "").strip()
        if email and email not in text:
            email = ""
        item = {
            "id": start_id + len(items),
            "name": str(record.get("name") or "").strip(),
            "school": str(record.get("school") or "").strip(),
            "college": str(record.get("college") or "").strip(),
            "title": str(record.get("title") or "").strip(),
            "direction": str(record.get("direction") or "").strip()[:180],
            "email": email,
            "admission_type": str(record.get("admission_type") or "").strip(),
            "source": str(record.get("source") or source).strip(),
            "raw_text": str(record.get("evidence") or "")[:1000],
            "match_score": 0,
            "match_level": "未评估",
            "match_reason": "",
            "risk_note": "",
        }
        key = (item["email"] or item["name"] + item["school"] + item["college"] + item["source"]).lower()
        if key in seen_keys or not _is_meaningful_supervisor(item):
            continue
        seen_keys.add(key)
        items.append(item)
    return items


def _llm_extract_supervisors_from_upload(text: str, source: str, start_id: int) -> list[dict[str, Any]]:
    settings = get_model_settings()
    if not settings.configured or not text.strip():
        return []
    try:
        from langchain_openai import ChatOpenAI
    except ImportError:
        return []

    prompt = f"""
你是 GradMate 的导师资料抽取器。请只根据用户上传文件中真实出现的文本抽取导师信息，不要猜测，不要编造。

信息来源：{source}

上传文件文本节选：
{text[:16000]}

抽取规则：
1. 只抽取明确的导师/教师/研究员个人信息。
2. 姓名、邮箱、职称、研究方向必须尽量属于同一个人；无法确认归属时对应字段留空。
3. 不要抽取学生、新闻作者、学院行政联系人、版权邮箱、招生办公共邮箱。
4. 每条记录必须有姓名，并且至少有邮箱、职称、研究方向、学院之一。
5. source 使用给定的信息来源。

只返回 JSON：
{{
  "supervisors": [
    {{
      "name": "姓名",
      "school": "学校",
      "college": "学院/院系",
      "title": "职称",
      "direction": "研究方向",
      "email": "邮箱",
      "admission_type": "招生类型",
      "source": "{source}",
      "evidence": "支持该记录的原文片段"
    }}
  ]
}}
"""
    try:
        model = ChatOpenAI(
            model=settings.model_id,
            api_key=settings.api_key,
            base_url=settings.base_url,
            temperature=0,
        )
        response = model.invoke(prompt)
        data = _extract_json_object(getattr(response, "content", str(response)))
    except Exception:
        return []

    records = data.get("supervisors", [])
    if not isinstance(records, list):
        return []

    items: list[dict[str, Any]] = []
    for record in records:
        if not isinstance(record, dict):
            continue
        email = str(record.get("email") or "").strip()
        if email and email not in text:
            email = ""
        item = {
            "id": start_id + len(items),
            "name": str(record.get("name") or "").strip(),
            "school": str(record.get("school") or "").strip(),
            "college": str(record.get("college") or "").strip(),
            "title": str(record.get("title") or "").strip(),
            "direction": str(record.get("direction") or "").strip()[:180],
            "email": email,
            "admission_type": str(record.get("admission_type") or "").strip(),
            "source": str(record.get("source") or source).strip(),
            "raw_text": str(record.get("evidence") or "")[:1000],
            "match_score": 0,
            "match_level": "未评估",
            "match_reason": "",
            "risk_note": "",
        }
        if _is_meaningful_supervisor(item):
            items.append(item)
    return items


def _near_email_chunks(text: str, max_chunks: int = 18) -> list[str]:
    lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
    email_line_indexes = [
        index for index, line in enumerate(lines)
        if re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", line)
    ]
    chunks: list[str] = []
    seen: set[str] = set()
    for index in email_line_indexes[:max_chunks]:
        start = max(0, index - 8)
        end = min(len(lines), index + 9)
        chunk = "\n".join(lines[start:end]).strip()
        key = re.sub(r"\s+", "", chunk)[:240]
        if chunk and key not in seen:
            chunks.append(chunk)
            seen.add(key)

    if chunks:
        return chunks

    window = 9000
    return [text[start:start + window] for start in range(0, min(len(text), window * 3), window) if text[start:start + window].strip()]


def _guess_name_from_chunk(chunk: str, email: str) -> str:
    lines = [line.strip() for line in chunk.splitlines() if line.strip()]
    email_index = next((index for index, line in enumerate(lines) if email in line), len(lines) - 1)
    candidates: list[str] = []
    for line in reversed(lines[max(0, email_index - 6): email_index + 1]):
        clean = re.sub(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", " ", line)
        clean = re.sub(r"(教授|副教授|讲师|研究员|副研究员|博士生导师|硕士生导师|邮箱|Email|E-mail|电话|职称|研究方向)", " ", clean, flags=re.I)
        candidates.extend(re.findall(r"[\u4e00-\u9fa5]{2,4}", clean))
    bad = {"邮箱", "电话", "主页", "学院", "大学", "研究", "方向", "教授", "导师", "中心", "团队", "办公室"}
    candidates = [item for item in candidates if item not in bad and not item.endswith(("学院", "大学"))]
    return candidates[-1] if candidates else ""


def _heuristic_extract_supervisors_from_upload(text: str, source: str, start_id: int) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    seen_emails: set[str] = set()
    title_pattern = r"(教授|副教授|讲师|研究员|副研究员|博士生导师|硕士生导师|博导|硕导)"
    direction_pattern = r"(?:研究方向|研究领域|主要研究|方向)\s*[:：]?\s*([^\n]{2,160})"
    for chunk in _near_email_chunks(text, max_chunks=40):
        emails = re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", chunk)
        for email in emails:
            if email in seen_emails:
                continue
            seen_emails.add(email)
            name = _guess_name_from_chunk(chunk, email)
            title_match = re.search(title_pattern, chunk)
            direction_match = re.search(direction_pattern, chunk)
            item = {
                "id": start_id + len(items),
                "name": name,
                "school": "",
                "college": "",
                "title": title_match.group(1) if title_match else "",
                "direction": direction_match.group(1).strip()[:180] if direction_match else "",
                "email": email,
                "admission_type": "",
                "source": source,
                "raw_text": chunk[:1000],
                "match_score": 0,
                "match_level": "未评估",
                "match_reason": "",
                "risk_note": "",
            }
            if _is_meaningful_supervisor(item):
                items.append(item)
    return items


def _llm_extract_supervisors_from_upload_chunks(text: str, source: str, start_id: int) -> list[dict[str, Any]]:
    settings = get_model_settings()
    if not settings.configured or not text.strip():
        return []
    try:
        from langchain_openai import ChatOpenAI
    except ImportError:
        return []

    model = ChatOpenAI(
        model=settings.model_id,
        api_key=settings.api_key,
        base_url=settings.base_url,
        temperature=0,
    )
    items: list[dict[str, Any]] = []
    seen_keys: set[str] = set()
    for chunk in _near_email_chunks(text, max_chunks=12):
        prompt = f"""
你是导师信息抽取器。只根据下面这一段 PDF 提取文本抽取导师/教师/研究员信息，不要猜测。

信息来源：{source}

文本片段：
{chunk[:6000]}

要求：
1. 只抽取真实个人教师信息，不抽取招生办、学院办公室、版权邮箱、学生、新闻作者。
2. 姓名和邮箱必须能在片段中找到，且尽量属于同一个人。
3. 没看到的字段留空。
4. 只返回 JSON：{{"supervisors":[{{"name":"","school":"","college":"","title":"","direction":"","email":"","admission_type":"","source":"{source}","evidence":""}}]}}
"""
        try:
            response = model.invoke(prompt)
            data = _extract_json_object(getattr(response, "content", str(response)))
        except Exception:
            continue
        records = data.get("supervisors", [])
        if not isinstance(records, list):
            continue
        for record in records:
            if not isinstance(record, dict):
                continue
            email = str(record.get("email") or "").strip()
            if email and email not in chunk:
                email = ""
            item = {
                "id": start_id + len(items),
                "name": str(record.get("name") or "").strip(),
                "school": str(record.get("school") or "").strip(),
                "college": str(record.get("college") or "").strip(),
                "title": str(record.get("title") or "").strip(),
                "direction": str(record.get("direction") or "").strip()[:180],
                "email": email,
                "admission_type": str(record.get("admission_type") or "").strip(),
                "source": str(record.get("source") or source).strip(),
                "raw_text": str(record.get("evidence") or chunk)[:1000],
                "match_score": 0,
                "match_level": "未评估",
                "match_reason": "",
                "risk_note": "",
            }
            key = (item["email"] or item["name"] + item["source"]).lower()
            if key in seen_keys or not _is_meaningful_supervisor(item):
                continue
            seen_keys.add(key)
            items.append(item)
    return items


def _llm_extract_supervisors_from_upload_direct(text: str, source: str, start_id: int) -> list[dict[str, Any]]:
    settings = get_model_settings()
    if not settings.configured or not text.strip():
        return []
    try:
        from langchain_openai import ChatOpenAI
    except ImportError:
        return []

    prompt = f"""
你是 GradMate 的导师资料实体抽取器。请直接从下面这段 PDF/文件提取文本中抽取导师、教师、研究员信息。

信息来源：{source}

文件文本：
{text[:24000]}

抽取要求：
1. 只抽取明确的个人导师、教师、研究员信息。
2. 不要抽取招生办公室、学院办公室、版权信息、学生、新闻作者、普通行政联系人。
3. 每条记录必须有姓名；邮箱可以为空。
4. 若文本中能看到职称、学院、研究方向、招生类型，请填入；看不到就留空，不要编造。
5. 如果 PDF 是导师列表，即使没有邮箱，也要抽取“姓名 + 职称/方向/学院”等有效记录。
6. evidence 填支持该导师的原文片段，source 固定使用给定信息来源。

只返回 JSON，不要解释：
{{
  "supervisors": [
    {{
      "name": "姓名",
      "school": "学校",
      "college": "学院或院系",
      "title": "职称",
      "direction": "研究方向",
      "email": "邮箱",
      "admission_type": "招生类型",
      "source": "{source}",
      "evidence": "原文片段"
    }}
  ]
}}
"""
    try:
        model = ChatOpenAI(
            model=settings.model_id,
            api_key=settings.api_key,
            base_url=settings.base_url,
            temperature=0,
        )
        response = model.invoke(prompt)
        data = _extract_json_object(getattr(response, "content", str(response)))
    except Exception as error:
        logger.warning("Direct LLM extraction failed: %s", error)
        return []
    return _normalize_llm_upload_records(data.get("supervisors", []), text, source, start_id)


def extract_supervisor_items(
    text: str,
    source: str,
    start_id: int,
    prefer_structured_table: bool = False,
) -> tuple[list[dict[str, Any]], str]:
    # Standard spreadsheets already have reliable field boundaries. Parse the
    # table first so an email-neighbour heuristic cannot merge adjacent rows.
    if prefer_structured_table:
        rule_items = [supervisor_to_dict(item) for item in extract_supervisors_from_text(text=text, source=source, start_id=start_id)]
        meaningful = [item for item in rule_items if _is_meaningful_supervisor(item)]
        if meaningful:
            return meaningful, "structured-table"

    direct_llm_items = _llm_extract_supervisors_from_upload_direct(text=text, source=source, start_id=start_id)
    if direct_llm_items:
        return direct_llm_items, "llm-direct"

    heuristic_items = _heuristic_extract_supervisors_from_upload(text=text, source=source, start_id=start_id)
    if heuristic_items:
        return heuristic_items, "email-context"

    llm_items = _llm_extract_supervisors_from_upload_chunks(text=text, source=source, start_id=start_id)
    if llm_items:
        return llm_items, "llm-chunks"

    rule_items = [supervisor_to_dict(item) for item in extract_supervisors_from_text(text=text, source=source, start_id=start_id)]
    meaningful = [item for item in rule_items if _is_meaningful_supervisor(item)]
    if meaningful:
        return meaningful, "rule"
    return [], "none"


def extract_upload_text(filename: str, content: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise HTTPException(status_code=500, detail="后端缺少 openpyxl，无法解析 XLSX。请安装 backend/requirements.txt。") from error

        import io

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    values = [re.sub(r"\s+", " ", str(value)).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        rows.append("\t".join(values))
                    if len(rows) >= 2000:
                        break
                if rows:
                    # Keep the header as the first line: the downstream table
                    # parser uses it to map Chinese/English spreadsheet fields.
                    if not lines:
                        lines.extend(rows)
                    elif rows[0] == lines[0]:
                        lines.extend(rows[1:])
                    else:
                        lines.extend(rows)
            workbook.close()
            return "\n".join(lines).strip()
        except Exception as error:
            raise HTTPException(status_code=400, detail=f"XLSX 解析失败：{error}") from error

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as error:
            raise HTTPException(status_code=500, detail="后端缺少 pypdf，无法解析 PDF。请安装 backend/requirements.txt。") from error

        import io

        try:
            reader = PdfReader(io.BytesIO(content))
            return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
        except Exception as error:
            raise HTTPException(status_code=400, detail=f"PDF 解析失败：{error}") from error

    for encoding in ["utf-8-sig", "utf-8", "gb18030"]:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise HTTPException(status_code=400, detail="暂不支持该文件编码或文件类型。")


def generate_question_suggestions(profile: Any, supervisor: dict[str, Any], focus: str = "") -> dict[str, Any]:
    settings = get_model_settings()
    if not settings.configured:
        raise HTTPException(status_code=400, detail="模型未配置，无法生成可能问题。")
    try:
        from langchain_openai import ChatOpenAI
    except ImportError as error:
        raise HTTPException(status_code=500, detail="缺少 langchain_openai，无法调用模型。") from error

    prompt = f"""
你是研究生申请面试辅导助手。请根据学生经历和目标导师方向，生成导师可能会问学生的问题。

学生画像：
- 姓名：{profile.name}
- 本科院校：{profile.university}
- 本科专业：{profile.major}
- 成绩排名：{profile.rank}
- 英语水平：{profile.english}
- 研究兴趣：{profile.interests}
- 项目经历：{profile.projects}
- 申请类型：{profile.application_type}
- 目标学校：{profile.target_schools}
- 目标学院：{profile.target_colleges}

导师资料：
- 姓名：{supervisor.get("name", "")}
- 学校：{supervisor.get("school", "")}
- 学院：{supervisor.get("college", "")}
- 职称：{supervisor.get("title", "")}
- 研究方向：{supervisor.get("direction", "")}
- 招生类型：{supervisor.get("admission_type", "")}
- 来源：{supervisor.get("source", "")}

用户关注点：{focus or "无"}

请生成 18-24 个问题，按类别组织：
1. 科研方向理解
2. 项目经历深挖
3. 专业基础与方法
4. 申请动机与匹配度
5. 风险追问
6. 学生可以反问导师的问题

要求：
- 问题必须贴合导师方向和学生经历。
- 不要写空泛模板问题。
- 每个问题给出“考察点”和“准备提示”。
- 只返回 JSON，不要解释。

JSON 格式：
{{
  "supervisor": "导师姓名",
  "summary": "一句话说明整体准备重点",
  "groups": [
    {{
      "title": "类别",
      "items": [
        {{"question": "问题", "intent": "考察点", "prep": "准备提示"}}
      ]
    }}
  ]
}}
"""
    model = ChatOpenAI(
        model=settings.model_id,
        api_key=settings.api_key,
        base_url=settings.base_url,
        temperature=0.2,
    )
    try:
        response = model.invoke(prompt)
    except Exception as error:
        raise HTTPException(status_code=500, detail=f"模型调用失败：{error}") from error

    data = _extract_json_object(getattr(response, "content", str(response)))
    groups = data.get("groups", [])
    if not isinstance(groups, list) or not groups:
        raise HTTPException(status_code=502, detail="模型未返回可用的问题结构。")
    return {
        "supervisor": data.get("supervisor") or supervisor.get("name", ""),
        "summary": data.get("summary") or "",
        "groups": groups,
    }


def _chunk_text(text: str, size: int = 900, overlap: int = 120) -> list[str]:
    clean = re.sub(r"\n{3,}", "\n\n", (text or "").strip())
    if not clean:
        return []
    chunks = []
    start = 0
    while start < len(clean):
        chunk = clean[start:start + size].strip()
        if chunk:
            chunks.append(chunk)
        start += max(1, size - overlap)
    return chunks


def _tokens_for_rag(text: str) -> set[str]:
    text = (text or "").lower()
    chinese_terms = re.findall(r"[\u4e00-\u9fa5]{2,}", text)
    ascii_terms = re.findall(r"[a-z0-9_+-]{2,}", text)
    stop = {"这个", "学院", "需要", "什么", "哪些", "是否", "要求", "报名", "导师"}
    return {item for item in chinese_terms + ascii_terms if item not in stop}


def _score_context(query: str, text: str) -> int:
    q_tokens = _tokens_for_rag(query)
    if not q_tokens:
        return 0
    compact_text = re.sub(r"\s+", "", (text or "").lower())
    score = 0
    for token in q_tokens:
        if token in compact_text:
            score += 3 if len(token) >= 4 else 1
    return score


def _knowledge_contexts(question: str, limit: int = 6) -> list[dict[str, Any]]:
    scored = []
    for item in get_knowledge_items():
        score = _score_context(question, item.get("text", ""))
        if score > 0:
            scored.append((score, item))
    scored.sort(key=lambda pair: pair[0], reverse=True)
    return [
        {
            "source": item.get("source", "知识库文件"),
            "kind": item.get("kind", "knowledge"),
            "snippet": item.get("text", "")[:900],
            "score": score,
        }
        for score, item in scored[:limit]
    ]


def _ddl_contexts(question: str, school: str = "", college: str = "", limit: int = 5) -> list[dict[str, Any]]:
    query = school or college or question[:20]
    contexts: list[dict[str, Any]] = []
    try:
        result = get_baoyan_ddl(query=query, limit=limit, refresh=False)
    except Exception:
        return contexts
    for item in result.get("items", []):
        snippet = "\n".join(
            [
                f"名称：{item.get('name', '')}",
                f"学院/单位：{item.get('institute', '')}",
                f"描述：{item.get('description', '')}",
                f"截止日期：{item.get('deadline_date', '') or item.get('deadline', '')}",
                f"链接：{item.get('website', '')}",
            ]
        )
        if _score_context(question + " " + school + " " + college, snippet) > 0:
            contexts.append({"source": item.get("website") or result.get("source_label", "保研DDL"), "kind": "ddl", "snippet": snippet})
    return contexts[:limit]


def _browser_contexts(question: str, school: str = "", college: str = "", limit: int = 3) -> list[dict[str, Any]]:
    query = " ".join(part for part in [school, college, question] if part).strip()
    if not query:
        return []
    try:
        from browser_research import _create_driver
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from urllib.parse import quote_plus
    except Exception:
        return []

    contexts: list[dict[str, Any]] = []
    driver = None
    try:
        driver = _create_driver()
        driver.get(f"https://www.bing.com/search?q={quote_plus(query)}")
        WebDriverWait(driver, 10).until(lambda item: item.execute_script("return document.readyState") in {"interactive", "complete"})
        links = []
        for element in driver.find_elements(By.CSS_SELECTOR, "li.b_algo h2 a, .b_algo a")[:8]:
            href = (element.get_attribute("href") or "").strip()
            title = (element.text or "").strip()
            if href.startswith(("http://", "https://")) and "bing.com" not in href:
                links.append((title, href))
        seen = set()
        for title, href in links:
            if len(contexts) >= limit or href in seen:
                continue
            seen.add(href)
            try:
                driver.get(href)
                WebDriverWait(driver, 8).until(lambda item: item.execute_script("return document.readyState") in {"interactive", "complete"})
                body = driver.find_element(By.TAG_NAME, "body").text
                contexts.append({"source": href, "kind": "browser", "snippet": f"{title}\n{body[:1200]}"})
            except Exception:
                continue
    except Exception as error:
        logger.warning("RAG browser lookup failed: %s", error)
    finally:
        if driver:
            driver.quit()
    return contexts


def _answer_with_rag(question: str, contexts: list[dict[str, Any]]) -> dict[str, Any]:
    if not contexts:
        return {"answer": "当前知识库和通知数据中没有找到足够相关的内容。", "sources": []}
    settings = get_model_settings()
    source_text = "\n\n".join(
        f"[{index + 1}] 来源：{ctx.get('source')}\n类型：{ctx.get('kind')}\n片段：{ctx.get('snippet')}"
        for index, ctx in enumerate(contexts)
    )
    if not settings.configured:
        return {
            "answer": "模型未配置，以下是检索到的相关来源片段，请人工核对后作答。",
            "sources": contexts,
        }
    try:
        from langchain_openai import ChatOpenAI
        model = ChatOpenAI(model=settings.model_id, api_key=settings.api_key, base_url=settings.base_url, temperature=0)
        prompt = f"""
你是 GradMate 的招生简章 RAG 问答助手。必须只根据给定来源片段回答，不能编造。

用户问题：{question}

来源片段：
{source_text}

回答要求：
1. 先直接回答问题。
2. 如果片段不足以确定，明确说“当前资料未找到明确说明”。
3. 每个关键结论后标注来源编号，例如 [1]。
4. 最后列出“依据来源”，包含文件名或链接和对应片段摘要。
"""
        response = model.invoke(prompt)
        answer = getattr(response, "content", str(response))
    except Exception as error:
        answer = f"模型调用失败，以下是检索到的相关来源片段：{error}"
    return {"answer": answer, "sources": contexts}


@asynccontextmanager
async def lifespan(_: FastAPI):
    ensure_data_dir()
    yield


app = FastAPI(
    title="GradMate API",
    description="导师匹配与邮件联系智能体后端，基于 FastAPI + LangChain create_agent。",
    version="0.2.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=get_security_settings().cors_origins,
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.include_router(email_router)
app.include_router(practice_router)
app.include_router(resume_router)

if FRONTEND_ASSETS_DIR.exists():
    app.mount("/assets", StaticFiles(directory=FRONTEND_ASSETS_DIR), name="assets")
else:
    app.mount("/static", StaticFiles(directory=FRONTEND_DIR), name="static")


@app.get("/")
def read_index() -> FileResponse:
    dist_index = FRONTEND_DIST_DIR / "index.html"
    if dist_index.exists():
        return FileResponse(dist_index)
    return FileResponse(FRONTEND_DIR / "index.html")


@app.get("/api/health")
def health() -> dict[str, Any]:
    config = agent_config()
    return {
        "ok": True,
        "message": "GradMate FastAPI backend is running",
        "agent_mode": config["agent_mode"],
        "agent_created": config["agent_created"],
    }


@app.get("/api/homepage/screenshot")
def homepage_screenshot(url: str = Query(..., min_length=8, max_length=1200)) -> dict[str, Any]:
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise HTTPException(status_code=400, detail="只支持 http/https 官网链接。")

    try:
        from selenium import webdriver
        from selenium.common.exceptions import WebDriverException
        from selenium.webdriver.chrome.options import Options as ChromeOptions
        from selenium.webdriver.edge.options import Options as EdgeOptions
    except ImportError as error:
        raise HTTPException(status_code=500, detail=f"后端缺少 selenium，无法截取官网预览：{error}") from error

    driver = None
    last_error: Exception | None = None
    browser_factories = []

    chrome_options = ChromeOptions()
    chrome_options.add_argument("--headless=new")
    chrome_options.add_argument("--disable-gpu")
    chrome_options.add_argument("--no-sandbox")
    chrome_options.add_argument("--disable-dev-shm-usage")
    chrome_options.add_argument("--hide-scrollbars")
    chrome_options.add_argument("--window-size=1280,820")
    browser_factories.append(lambda: webdriver.Chrome(options=chrome_options))

    edge_options = EdgeOptions()
    edge_options.add_argument("--headless=new")
    edge_options.add_argument("--disable-gpu")
    edge_options.add_argument("--window-size=1280,820")
    browser_factories.append(lambda: webdriver.Edge(options=edge_options))

    for factory in browser_factories:
        try:
            driver = factory()
            break
        except WebDriverException as error:
            last_error = error

    if driver is None:
        raise HTTPException(status_code=503, detail=f"没有可用的浏览器驱动，无法生成官网预览：{last_error}")

    try:
        driver.set_page_load_timeout(14)
        driver.set_script_timeout(8)
        driver.get(url.strip())
        png = driver.get_screenshot_as_png()
        return {
            "url": url.strip(),
            "image": f"data:image/png;base64,{base64.b64encode(png).decode('ascii')}",
        }
    except Exception as error:
        raise HTTPException(status_code=502, detail=f"官网预览截图失败：{error}") from error
    finally:
        driver.quit()


@app.get("/api/dashboard")
def dashboard_summary() -> dict[str, Any]:
    supervisors = get_supervisors()
    high_match = sum(1 for item in supervisors if int(item.get("match_score") or 0) >= 85)
    missing_email = sum(1 for item in supervisors if not item.get("email"))
    draft_count = sum(1 for item in supervisors if item.get("email_body"))
    return {
        "supervisor_count": len(supervisors),
        "high_match_count": high_match,
        "missing_email_count": missing_email,
        "draft_count": draft_count,
    }


@app.get("/api/baoyan-ddl")
def baoyan_ddl(
    source: str = "camp2026",
    query: str = "",
    limit: int = 40,
    refresh: bool = False,
) -> dict[str, Any]:
    return get_baoyan_ddl(source=source, query=query, limit=limit, refresh=refresh)


def _validate_report_format(report_format: str) -> str:
    value = (report_format or "docx").lower()
    if value not in {"docx", "xlsx"}:
        raise HTTPException(status_code=400, detail="format must be docx or xlsx")
    return value


def _download_report(path: Path) -> FileResponse:
    media_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if path.suffix.lower() == ".xlsx"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    return FileResponse(path, media_type=media_type, filename=path.name)


@app.get("/api/export/supervisors")
def export_supervisors(format: str = "docx") -> FileResponse:
    report_format = _validate_report_format(format)
    path = export_supervisors_report(get_supervisors(), report_format=report_format)
    return _download_report(path)


@app.get("/api/export/matches")
def export_matches(format: str = "docx") -> FileResponse:
    report_format = _validate_report_format(format)
    supervisors = sorted(get_supervisors(), key=lambda item: int(item.get("match_score") or 0), reverse=True)
    path = export_supervisors_report(supervisors, report_format=report_format, title="GradMate Supervisor Match Report")
    return _download_report(path)


@app.get("/api/export/ddl")
def export_ddl(source: str = "camp2026", query: str = "", limit: int = 80, format: str = "docx") -> FileResponse:
    report_format = _validate_report_format(format)
    data = get_baoyan_ddl(source=source, query=query, limit=limit, refresh=False)
    path = export_ddl_report(data.get("items", []), report_format=report_format)
    return _download_report(path)


@app.get("/api/export/knowledge")
def export_knowledge(format: str = "docx") -> FileResponse:
    report_format = _validate_report_format(format)
    path = export_knowledge_report(get_knowledge_items(), report_format=report_format)
    return _download_report(path)


@app.post("/api/export/rag-answer")
def export_rag_answer(payload: ExportRagAnswerRequest, format: str = "docx") -> FileResponse:
    report_format = _validate_report_format(format)
    path = export_rag_answer_report(payload.dict(), report_format=report_format)
    return _download_report(path)


@app.post("/api/export/questions")
def export_questions(payload: ExportQuestionsRequest, format: str = "docx") -> FileResponse:
    report_format = _validate_report_format(format)
    path = export_questions_report(payload.dict(), report_format=report_format)
    return _download_report(path)


@app.get("/api/profile")
def read_profile() -> dict[str, Any]:
    profile = get_profile()
    return {"profile": profile.__dict__, "summary": profile_summary(profile)}


@app.post("/api/profile")
def write_profile(payload: dict[str, Any]) -> dict[str, Any]:
    profile = profile_from_dict(payload)
    save_profile(profile)
    return {"profile": profile.__dict__, "summary": profile_summary(profile)}


@app.get("/api/supervisors")
def read_supervisors() -> dict[str, Any]:
    return {"supervisors": get_supervisors()}


@app.delete("/api/supervisors/{supervisor_id}")
def delete_supervisor(supervisor_id: int) -> dict[str, Any]:
    current = get_supervisors()
    kept = [item for item in current if int(item.get("id", 0) or 0) != supervisor_id]
    if len(kept) == len(current):
        raise HTTPException(status_code=404, detail="Supervisor not found")
    save_supervisors(kept)
    return {"deleted": supervisor_id, "supervisors": kept}


@app.patch("/api/supervisors/{supervisor_id}/contact-status")
def update_supervisor_contact_status(supervisor_id: int, payload: SupervisorContactStatusRequest) -> dict[str, Any]:
    manual_statuses = {"等待回复", "已回复", "需补材料", "已约面谈", "暂不招生"}
    status = payload.status.strip()
    if status not in manual_statuses:
        raise HTTPException(status_code=400, detail="该状态由系统自动跟进，不能手动设置。")

    supervisors = get_supervisors()
    for item in supervisors:
        if int(item.get("id", 0) or 0) == supervisor_id:
            item["contact_status"] = status
            item["contact_status_mode"] = "manual"
            item["contact_last_at"] = date.today().isoformat()
            save_supervisors(supervisors)
            return {"supervisor": item, "supervisors": supervisors}
    raise HTTPException(status_code=404, detail="Supervisor not found")


@app.get("/api/school-directory")
def read_school_directory(school: str = "", major: str = "", college: str = "") -> dict[str, Any]:
    if school or major or college:
        matches = find_school_sites(school=school, major=major, college=college)
        return {
            "count": len(matches),
            "items": [
                {
                    "school": item.school,
                    "department": item.department,
                    "major": item.major,
                    "url": item.url,
                    "source_sheet": item.source_sheet,
                }
                for item in matches
            ],
        }
    return school_directory_summary()


@app.post("/api/supervisors/upload")
def upload_supervisors(payload: SupervisorUploadRequest) -> dict[str, Any]:
    current = get_supervisors()
    start_id = max([int(item.get("id", 0)) for item in current] + [0]) + 1
    items, method = extract_supervisor_items(text=payload.text, source=payload.source, start_id=start_id)
    supervisors, added = merge_supervisors(current, items)
    save_supervisors(supervisors)
    return {"imported": len(added), "supervisors": supervisors, "method": method, "text_length": len(payload.text or "")}


@app.post("/api/supervisors/upload-file")
async def upload_supervisor_file(
    file: UploadFile = File(...),
    source: str = Form("用户上传资料"),
) -> dict[str, Any]:
    content = await file.read()
    text = extract_upload_text(file.filename or "", content)
    if not text.strip():
        raise HTTPException(status_code=400, detail="文件中没有提取到可用文本。")

    current = get_supervisors()
    start_id = max([int(item.get("id", 0)) for item in current] + [0]) + 1
    items, method = extract_supervisor_items(
        text=text,
        source=source or file.filename or "用户上传资料",
        start_id=start_id,
        prefer_structured_table=Path(file.filename or "").suffix.lower() == ".xlsx",
    )
    supervisors, added = merge_supervisors(current, items)
    save_supervisors(supervisors)
    logger.info(
        "Supervisor upload: file=%s text_length=%s method=%s extracted=%s imported=%s",
        file.filename,
        len(text),
        method,
        len(items),
        len(added),
    )
    return {"imported": len(added), "supervisors": supervisors, "text_length": len(text), "method": method}


@app.post("/api/supervisors/match")
def match_supervisors() -> dict[str, Any]:
    profile = get_profile()
    scored = [
        supervisor_to_dict(item)
        for item in score_supervisors_hybrid(
            profile,
            [supervisor_from_dict(item) for item in get_supervisors()],
        )
    ]
    save_supervisors(scored)
    return {"supervisors": scored, "profile_summary": profile_summary(profile)}


@app.get("/api/knowledge")
def read_knowledge() -> dict[str, Any]:
    items = get_knowledge_items()
    files: dict[str, int] = {}
    for item in items:
        source = item.get("source", "知识库文件")
        files[source] = files.get(source, 0) + 1
    return {"chunk_count": len(items), "files": [{"source": key, "chunks": value} for key, value in files.items()]}


@app.get("/api/rag/history")
def read_rag_history(limit: int = 30) -> dict[str, Any]:
    safe_limit = min(max(limit, 1), 100)
    return {"items": get_rag_history()[:safe_limit]}


@app.get("/api/questions/history")
def read_question_history(limit: int = 30) -> dict[str, Any]:
    safe_limit = min(max(limit, 1), 100)
    return {"items": get_question_history()[:safe_limit]}


@app.post("/api/knowledge/upload")
async def upload_knowledge_file(
    file: UploadFile = File(...),
    kind: str = Form("招生简章"),
) -> dict[str, Any]:
    content = await file.read()
    text = extract_upload_text(file.filename or "", content)
    chunks = _chunk_text(text)
    if not chunks:
        raise HTTPException(status_code=400, detail="文件中没有提取到可用文本。")
    current = get_knowledge_items()
    source = file.filename or "上传文件"
    current = [item for item in current if item.get("source") != source]
    start_id = max([int(item.get("id", 0) or 0) for item in current] + [0]) + 1
    for offset, chunk in enumerate(chunks):
        current.append({"id": start_id + offset, "source": source, "kind": kind, "text": chunk})
    save_knowledge_items(current)
    return {"source": source, "kind": kind, "chunks": len(chunks), "text_length": len(text), "chunk_count": len(current)}


@app.post("/api/rag/ask")
def ask_rag(payload: RagAskRequest) -> dict[str, Any]:
    question = payload.question.strip()
    if not question:
        raise HTTPException(status_code=400, detail="问题不能为空。")
    contexts = _knowledge_contexts(question)
    ddl_contexts = _ddl_contexts(question, school=payload.school, college=payload.college)
    contexts.extend(ddl_contexts)
    used_browser = False
    if payload.use_browser and len(contexts) < 2:
        browser_items = _browser_contexts(question, school=payload.school, college=payload.college)
        contexts.extend(browser_items)
        used_browser = bool(browser_items)
    answer = _answer_with_rag(question, contexts[:10])
    result = {
        "answer": answer["answer"],
        "sources": answer["sources"],
        "used_browser": used_browser,
        "context_count": len(contexts),
    }
    history = get_rag_history()
    history.insert(0, {
        "id": datetime.now().strftime("%Y%m%d%H%M%S%f"),
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "question": question,
        "school": payload.school.strip(),
        "college": payload.college.strip(),
        **result,
    })
    save_rag_history(history[:100])
    return result


@app.post("/api/questions/generate")
def generate_questions(payload: QuestionGenerateRequest) -> dict[str, Any]:
    profile = get_profile()
    for item in get_supervisors():
        if int(item.get("id", 0) or 0) == payload.supervisor_id:
            result = generate_question_suggestions(profile, item, payload.focus)
            response = {"questions": result, "profile_summary": profile_summary(profile)}
            history = get_question_history()
            history.insert(0, {
                "id": datetime.now().strftime("%Y%m%d%H%M%S%f"),
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "supervisor_id": payload.supervisor_id,
                "supervisor_name": item.get("name", ""),
                "focus": payload.focus,
                "questions": result,
            })
            save_question_history(history[:100])
            return response
    raise HTTPException(status_code=404, detail="Supervisor not found")


@app.post("/api/supervisors/agent-browse")
def browse_supervisors(payload: SupervisorBrowseRequest) -> dict[str, Any]:
    try:
        return browse_school_supervisors(
            school=payload.school,
            major=payload.major,
            college=payload.college,
            keywords=payload.keywords,
            max_pages=payload.max_pages,
            save=payload.save,
        )
    except RuntimeError as error:
        raise HTTPException(status_code=409, detail=str(error)) from error


@app.get("/api/agent/config")
def read_agent_config() -> dict[str, Any]:
    return agent_config()


@app.post("/api/agent/invoke")
def run_agent(payload: AgentInvokeRequest) -> dict[str, Any]:
    try:
        return invoke_agent(payload.message)
    except RuntimeError as error:
        raise HTTPException(status_code=503, detail=str(error)) from error


@app.post("/api/assistant/run")
def run_assistant_task(payload: AssistantTaskRequest) -> dict[str, Any]:
    return run_application_assistant(payload.message)


@app.get("/api/assistant/history")
def read_assistant_history(limit: int = Query(default=30, ge=1, le=100)) -> dict[str, Any]:
    return {"items": get_assistant_history(limit)}


@app.post("/api/assistant/stream")
async def stream_assistant_task(payload: AssistantTaskRequest) -> StreamingResponse:
    """Stream real controller events so the UI shows the tool chain as it runs."""
    loop = asyncio.get_running_loop()
    event_queue: asyncio.Queue[dict[str, Any]] = asyncio.Queue()

    def publish(event: dict[str, Any]) -> None:
        # Wait until the event has entered the async queue so ordering remains
        # faithful to the actual tool execution order.
        asyncio.run_coroutine_threadsafe(
            event_queue.put({"type": "event", "data": event}), loop
        ).result()

    async def run_controller() -> None:
        try:
            result = await asyncio.to_thread(run_application_assistant, payload.message, publish)
            await event_queue.put({"type": "result", "data": result})
        except Exception as error:
            await event_queue.put({"type": "error", "data": {"detail": str(error)}})

    async def event_stream():
        controller_task = asyncio.create_task(run_controller())
        while True:
            item = await event_queue.get()
            yield f"event: {item['type']}\ndata: {json.dumps(item['data'], ensure_ascii=False)}\n\n"
            if item["type"] in {"result", "error"}:
                break
        await controller_task

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


if __name__ == "__main__":
    import uvicorn

    server = get_server_settings()
    ensure_data_dir()
    uvicorn.run("app:app", host=server.host, port=server.port, reload=False)
